"""
generar_fichas_guadalajara.py
Genera un Excel tipo ficha Fluix para cada negocio de estética/medicina
estética de Guadalajara extraído de mejoresclinicas.net
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

OUTPUT_DIR = r"C:\Users\Samu\Downloads\fichas_guadalajara"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ─── Estilos ────────────────────────────────────────────────────────────────
HDR_FILL   = PatternFill("solid", fgColor="1F3864")
HDR_FONT   = Font(bold=True, color="FFFFFF", size=11)
LABEL_FILL = PatternFill("solid", fgColor="D6E4F7")
LABEL_FONT = Font(bold=True, color="1F3864", size=10)
NOTE_FILL  = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT  = Font(color="7F6000", size=9, italic=True)
BODY_FONT  = Font(size=10)
THIN       = Side(style="thin", color="BFBFBF")
BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP       = Alignment(wrap_text=True, vertical="top")
CENTER     = Alignment(horizontal="center", vertical="center", wrap_text=True)


def hdr(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CENTER, BORDER


def label(ws, row, text):
    c = ws.cell(row=row, column=1, value=text)
    c.fill, c.font, c.alignment, c.border = LABEL_FILL, LABEL_FONT, WRAP, BORDER


def val(ws, row, text):
    c = ws.cell(row=row, column=2, value=text)
    c.font, c.alignment, c.border = BODY_FONT, WRAP, BORDER


def note(ws, row, text):
    c = ws.cell(row=row, column=3, value=text)
    c.fill, c.font, c.alignment, c.border = NOTE_FILL, NOTE_FONT, WRAP, BORDER


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_info_sheet(wb, b):
    ws = wb.active
    ws.title = "Info"
    ws.row_dimensions[1].height = 22
    hdr(ws, 1, 1, "CAMPO"); hdr(ws, 1, 2, "VALOR"); hdr(ws, 1, 3, "ESTADO / NOTAS")
    rows = [
        ("Nombre del negocio",  b["nombre"],        "Confirmado en directorio"),
        ("Slogan / Tagline",    b.get("slogan",""),  "Propuesta; confirmar con el negocio"),
        ("Descripción",         b["descripcion"],    "Redacción a partir de directorio"),
        ("Categoría",           b["categoria"],      "Enum CategoriaNegocio"),
        ("Precio medio",        b.get("precio",""),  "No publicado de forma verificable"),
        ("Especialidades",      b["especialidades"], "Servicios anunciados en directorio"),
        ("Rating de Google",    f"{b['rating']}/5 ({b['resenas']} reseñas, fuente directorio)", "Verificar en Google antes de guardar"),
        ("Número de reseñas",   str(b["resenas"]),   "Dato de directorio externo; puede variar"),
        ("Google Place ID",     "",                  "No localizado/verificado"),
    ]
    for i, (campo, valor, estado) in enumerate(rows, 2):
        label(ws, i, campo); val(ws, i, valor); note(ws, i, estado)
        ws.row_dimensions[i].height = 36 if "Descripción" in campo or "Especialidades" in campo else 20
    set_col_widths(ws, [28, 55, 45])


def add_contacto_sheet(wb, b):
    ws = wb.create_sheet("Contacto")
    hdr(ws, 1, 1, "CAMPO"); hdr(ws, 1, 2, "VALOR"); hdr(ws, 1, 3, "ESTADO / NOTAS")
    maps_url = (
        "https://www.google.com/maps/search/?api=1&query="
        + b["direccion"].replace(" ", "+").replace(",", "%2C")
    )
    rows = [
        ("Dirección",              b["direccion"],          "Confirmado"),
        ("Teléfono",               b["telefono"],           "Confirmado en directorio"),
        ("WhatsApp",               b.get("whatsapp",""),    "Verificar si es número de negocio"),
        ("Email público",          b.get("email",""),       "No publicado/verificado" if not b.get("email") else "Confirmado"),
        ("Email de notificaciones",b.get("email",""),       "Propuesta provisional"),
        ("Página web",             b.get("web",""),         "Confirmar que está activa"),
        ("Instagram",              b.get("instagram",""),   "No localizado/verificado" if not b.get("instagram") else "Confirmar"),
        ("Facebook",               b.get("facebook",""),   "No localizado/verificado" if not b.get("facebook") else "Confirmar"),
        ("URL de Google Maps",     maps_url,                "Enlace de búsqueda por dirección"),
    ]
    for i, (campo, valor, estado) in enumerate(rows, 2):
        label(ws, i, campo); val(ws, i, valor); note(ws, i, estado)
        ws.row_dimensions[i].height = 20
    set_col_widths(ws, [28, 55, 45])


def add_horario_sheet(wb, b):
    ws = wb.create_sheet("Horario")
    cols = ["DÍA","ABIERTO","FRANJA 1 APERTURA","FRANJA 1 CIERRE","FRANJA 2 APERTURA","FRANJA 2 CIERRE","NOTAS"]
    for j, c in enumerate(cols, 1):
        hdr(ws, 1, j, c)
    for i, fila in enumerate(b["horario"], 2):
        for j, v in enumerate(fila, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, BORDER
        ws.row_dimensions[i].height = 18
    set_col_widths(ws, [14, 10, 20, 20, 20, 20, 38])


def add_servicios_sheet(wb, b):
    ws = wb.create_sheet("Servicios")
    cols = ["NOMBRE DEL SERVICIO","DESCRIPCIÓN","CATEGORÍA","PRECIO","MODALIDAD PRECIO",
            "DURACIÓN MIN","PÚBLICO","ACTIVO","FOTO","ESTADO / NOTAS"]
    for j, c in enumerate(cols, 1):
        hdr(ws, 1, j, c)
    nota = "Servicio anunciado en directorio; precio y duración no publicados/verificados"
    for i, srv in enumerate(b["servicios"], 2):
        valores = [srv.get("nombre",""), "", srv.get("categoria",""), None, "Por definir", None, "General", "Sí", None, nota]
        for j, v in enumerate(valores, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font   = NOTE_FONT if j == 10 else BODY_FONT
            cell.fill   = NOTE_FILL if j == 10 else PatternFill()
            cell.alignment, cell.border = WRAP, BORDER
        ws.row_dimensions[i].height = 18
    set_col_widths(ws, [32, 28, 22, 10, 16, 14, 10, 8, 8, 48])


def add_extras_sheet(wb, b):
    ws = wb.create_sheet("Extras")
    hdr(ws, 1, 1, "CAMPO"); hdr(ws, 1, 2, "VALOR"); hdr(ws, 1, 3, "ESTADO / NOTAS")
    rows = [
        ("Empresa vinculada",          "",                   "Seleccionar desde Fluix"),
        ("Negocio destacado",          "No definido",        "Decisión interna"),
        ("Reservas online activadas",  "Por verificar",      "Comprobar antes de activar"),
        ("Acepta tarjeta",             b.get("tarjeta",""),  "No verificado"),
        ("Parking",                    "",                   "No verificado"),
        ("Accesible en silla de ruedas","",                  "No verificado"),
        ("WiFi",                       "",                   "No verificado"),
        ("Admite mascotas",            "",                   "No verificado"),
        ("Terraza",                    "No aplica",          "No verificado"),
        ("Eliminar negocio",           "No ejecutar",        "Acción destructiva"),
    ]
    for i, (campo, valor, estado) in enumerate(rows, 2):
        label(ws, i, campo); val(ws, i, valor); note(ws, i, estado)
        ws.row_dimensions[i].height = 18
    set_col_widths(ws, [28, 30, 45])


def add_formulario_sheet(wb, b):
    ws = wb.create_sheet("Formulario")
    hdr(ws, 1, 1, "CAMPO"); hdr(ws, 1, 2, "VALOR SUGERIDO"); hdr(ws, 1, 3, "ESTADO / NOTAS")
    rows = [
        ("Título del formulario",  f"Pide tu cita en {b['nombre']}",  "Adaptar al CTA del negocio"),
        ("Texto del botón",        "Pedir cita",                       ""),
        ("Nombre",                 "Texto; obligatorio",               ""),
        ("Teléfono",               "Teléfono; obligatorio",            ""),
        ("Email",                  "Email; opcional",                  ""),
        ("Servicio de interés",    "Desplegable; opcional",            ""),
        ("Fecha preferida",        "Fecha; opcional",                  ""),
        ("Hora preferida",         "Hora; opcional",                   ""),
        ("Mensaje",                "Párrafo; opcional",                ""),
    ]
    for i, (campo, valor, estado) in enumerate(rows, 2):
        label(ws, i, campo); val(ws, i, valor); note(ws, i, estado)
        ws.row_dimensions[i].height = 18
    set_col_widths(ws, [28, 35, 45])


def add_fuentes_sheet(wb, b):
    ws = wb.create_sheet("Fuentes")
    hdr(ws, 1, 1, "FUENTE"); hdr(ws, 1, 2, "URL"); hdr(ws, 1, 3, "INFORMACIÓN EXTRAÍDA")
    rows = [
        ("Directorio mejoresclinicas.net", b["fuente_url"],
         "Nombre, dirección, teléfono, horario, rating, reseñas, servicios"),
    ]
    if b.get("web"):
        rows.append(("Web oficial (no verificada)", b["web"], "Pendiente de verificar"))
    rows.append(("PENDIENTE DE VERIFICAR", "",
                 "Web oficial, Google Place ID, email, redes sociales, precios, accesibilidad"))
    for i, (fuente, url, info) in enumerate(rows, 2):
        label(ws, i, fuente); val(ws, i, url); note(ws, i, info)
        ws.row_dimensions[i].height = 20
    set_col_widths(ws, [30, 60, 50])


def crear_excel(b):
    wb = openpyxl.Workbook()
    add_info_sheet(wb, b)
    add_contacto_sheet(wb, b)
    add_horario_sheet(wb, b)
    add_servicios_sheet(wb, b)
    add_extras_sheet(wb, b)
    add_formulario_sheet(wb, b)
    add_fuentes_sheet(wb, b)
    slug = b["nombre"].lower().replace(" ", "_").replace("/","").replace("á","a").replace("é","e").replace("í","i").replace("ó","o").replace("ú","u").replace("ñ","n").replace(".","").replace(",","")
    path = os.path.join(OUTPUT_DIR, f"{slug}_ficha_fluix.xlsx")
    wb.save(path)
    print(f"OK {b['nombre']}  ->  {path}")


# ─── Datos de los 14 negocios ────────────────────────────────────────────────
NEGOCIOS = [
    {
        "nombre": "Belen Nogueras Estética",
        "slogan": "Los mejores cuidados naturales para tu belleza",
        "descripcion": "Centro especializado en estética facial y corporal ubicado en Aguas Vivas, Guadalajara. Destacado por el uso de productos ecológicos y naturales y por su trato personalizado.",
        "categoria": "Estéticas",
        "especialidades": "Tratamientos faciales, limpiezas faciales, depilación con cera, diseño de cejas, manicura, pedicura, masajes RITUAL SARTORI, asesoramiento en rutinas de cuidado de piel",
        "precio": "",
        "rating": "5.0",
        "resenas": 117,
        "direccion": "Av. de Beleña, 10, 19005 Guadalajara",
        "telefono": "722 26 27 08",
        "whatsapp": "722 26 27 08",
        "email": "",
        "web": "https://www.belennoguerasestetica.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "14:00", "16:00", "19:00", ""),
            ("Martes",    "Sí", "10:00", "14:00", "16:00", "19:00", ""),
            ("Miércoles", "Sí", "10:00", "14:00", "16:00", "19:00", ""),
            ("Jueves",    "Sí", "10:00", "14:00", "16:00", "19:00", ""),
            ("Viernes",   "Sí", "10:00", "14:00", "16:00", "19:00", ""),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Tratamientos faciales",                "categoria": "Belleza facial"},
            {"nombre": "Limpiezas faciales",                   "categoria": "Belleza facial"},
            {"nombre": "Diseño y depilación de cejas",         "categoria": "Diseño de la mirada"},
            {"nombre": "Depilación con cera",                  "categoria": "Depilación"},
            {"nombre": "Manicura",                             "categoria": "Estilo personal"},
            {"nombre": "Pedicura",                             "categoria": "Estilo personal"},
            {"nombre": "Masajes RITUAL SARTORI",               "categoria": "Bienestar"},
            {"nombre": "Asesoramiento en rutinas de cuidado",  "categoria": "Belleza facial"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/belen-nogueras-estetica/",
    },
    {
        "nombre": "Clínica Médico Estética Cristina Aguilera",
        "slogan": "Medicina estética con resultados naturales",
        "descripcion": "Centro especializado en medicina estética facial ubicado en el barrio Centro de Guadalajara. Uno de los centros mejor valorados de la zona según directorio.",
        "categoria": "Medicina estética",
        "especialidades": "Limpieza profunda facial, tratamientos antiinflamatorios, masajes estéticos, retoques estéticos, medicina estética facial general",
        "precio": "",
        "rating": "5.0",
        "resenas": 41,
        "direccion": "Pl. de Pablo Iglesias, 4, 19001 Guadalajara",
        "telefono": "659 78 09 80",
        "whatsapp": "659 78 09 80",
        "email": "",
        "web": "https://www.cristinaguileraesthetic.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "14:00", "16:00", "20:00", ""),
            ("Martes",    "Sí", "10:00", "14:00", "16:00", "20:00", ""),
            ("Miércoles", "Sí", "10:00", "14:00", "16:00", "20:00", ""),
            ("Jueves",    "Sí", "10:00", "14:00", "16:00", "20:00", ""),
            ("Viernes",   "Sí", "10:00", "14:00", "16:00", "20:00", ""),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Limpieza profunda facial",          "categoria": "Belleza facial"},
            {"nombre": "Tratamientos antiinflamatorios",    "categoria": "Belleza facial"},
            {"nombre": "Masajes estéticos",                 "categoria": "Bienestar"},
            {"nombre": "Retoques estéticos",                "categoria": "Medicina estética"},
            {"nombre": "Medicina estética facial",          "categoria": "Medicina estética"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/clinica-medico-estetica-cristina-aguilera/",
    },
    {
        "nombre": "Sensorial Estética Avanzada",
        "slogan": "Estética avanzada con trato cercano",
        "descripcion": "Centro especializado en estética avanzada ubicado en el barrio de Amistad, Guadalajara. Destaca por sus tratamientos antiedad y su trato profesional y cercano.",
        "categoria": "Estéticas",
        "especialidades": "Tratamientos faciales antiedad (arrugas, flacidez), bio lifting Nazarov, cavitación, láser, limpiezas faciales, línea de productos cosméticos",
        "precio": "",
        "rating": "5.0",
        "resenas": 125,
        "direccion": "Av. de Castilla, 18, G, 19002 Guadalajara",
        "telefono": "613 17 79 18",
        "whatsapp": "613 17 79 18",
        "email": "",
        "web": "https://www.sensorialesteticaavanzada.com",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "20:00", "", "", "Sin pausa al mediodía"),
            ("Martes",    "Sí", "10:00", "20:00", "", "", ""),
            ("Miércoles", "Sí", "10:00", "20:00", "", "", ""),
            ("Jueves",    "Sí", "10:00", "20:00", "", "", ""),
            ("Viernes",   "Sí", "10:00", "20:00", "", "", ""),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Tratamientos faciales antiedad",   "categoria": "Belleza facial"},
            {"nombre": "Bio lifting Nazarov",              "categoria": "Belleza facial"},
            {"nombre": "Cavitación",                       "categoria": "Belleza corporal"},
            {"nombre": "Láser estético",                   "categoria": "Tratamientos avanzados"},
            {"nombre": "Limpieza facial",                  "categoria": "Belleza facial"},
            {"nombre": "Línea de productos cosméticos",    "categoria": "Productos"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/sensorial-estetica-avanzada/",
    },
    {
        "nombre": "Clínica Medicina Estética Dra. Ana Sánchez Domínguez",
        "slogan": "Medicina estética con técnicas novedosas",
        "descripcion": "Clínica de medicina estética facial ubicada en el barrio Centro de Guadalajara. Dirigida por la Dra. Ana Sánchez Domínguez, especializada en técnicas novedosas y actualizadas.",
        "categoria": "Medicina estética",
        "especialidades": "Medicina estética facial, técnicas novedosas y actualizadas",
        "precio": "",
        "rating": "4.9",
        "resenas": 32,
        "direccion": "C. Ángel Martín Puebla, 5, local A, 19001 Guadalajara",
        "telefono": "638 05 28 11",
        "whatsapp": "638 05 28 11",
        "email": "",
        "web": "https://www.doctoraanasanchez.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "14:00", "",      "",      "Mañana solo"),
            ("Martes",    "Sí", "16:00", "20:00", "",      "",      "Tarde solo"),
            ("Miércoles", "Sí", "10:00", "16:00", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:00", "14:00", "17:00", "20:00", ""),
            ("Viernes",   "Sí", "10:00", "14:00", "",      "",      "Mañana solo"),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Medicina estética facial",           "categoria": "Medicina estética"},
            {"nombre": "Técnicas de rejuvenecimiento facial","categoria": "Medicina estética"},
            {"nombre": "Diagnóstico / valoración inicial",   "categoria": "Diagnóstico"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/clinica-medicina-estetica-dra-ana-sanchez-dominguez/",
    },
    {
        "nombre": "Anatura Bioestética",
        "slogan": "Estética y bienestar, también para pacientes oncológicos",
        "descripcion": "Centro de estética ubicado en El Fuerte, Guadalajara. Único en la zona en asesoramiento y venta de pelucas y turbantes oncológicos. Ofrece tratamientos faciales, masajes y depilación.",
        "categoria": "Estéticas",
        "especialidades": "Tratamientos faciales, exfoliación con punta de diamante, masajes, depilación, asesoramiento oncológico, venta de pelucas y turbantes para pacientes de cáncer",
        "precio": "",
        "rating": "4.9",
        "resenas": 40,
        "direccion": "C. del Arrabal del Agua, 1, local 2, 19005 Guadalajara",
        "telefono": "653 15 57 74",
        "whatsapp": "653 15 57 74",
        "email": "",
        "web": "",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "13:00", "16:30", "19:00", ""),
            ("Martes",    "Sí", "10:30", "13:00", "16:30", "19:00", ""),
            ("Miércoles", "Sí", "10:00", "13:00", "16:30", "19:00", ""),
            ("Jueves",    "Sí", "10:30", "13:00", "16:30", "19:00", ""),
            ("Viernes",   "Sí", "10:00", "13:00", "16:30", "19:00", ""),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Tratamientos faciales",                       "categoria": "Belleza facial"},
            {"nombre": "Exfoliación con punta de diamante",           "categoria": "Belleza facial"},
            {"nombre": "Masajes",                                     "categoria": "Bienestar"},
            {"nombre": "Depilación",                                   "categoria": "Depilación"},
            {"nombre": "Asesoramiento pelucas oncológicas",           "categoria": "Oncología estética"},
            {"nombre": "Venta de pelucas y turbantes oncológicos",    "categoria": "Oncología estética"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/anatura-bioestetica-centro-de-estetica-en-guadalajara-venta-de-pelucas-en-guadalajara/",
    },
    {
        "nombre": "Beauty by Patricia",
        "slogan": "Tu centro de belleza en Entrepeñas",
        "descripcion": "Centro especializado en medicina estética facial ubicado en Aguas Vivas, Guadalajara. Uno de los centros mejor valorados de la zona según directorio.",
        "categoria": "Estéticas",
        "especialidades": "Depilación láser, manicura, tratamientos faciales, limpiezas faciales, medicina estética facial",
        "precio": "",
        "rating": "4.9",
        "resenas": 97,
        "direccion": "Blvr. de Entrepeñas, 1, B, 19005 Guadalajara",
        "telefono": "673 81 17 68",
        "whatsapp": "673 81 17 68",
        "email": "",
        "web": "https://www.beautyguadalajara.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "15:30", "18:00", "",      "",      "Solo tarde"),
            ("Martes",    "Sí", "10:00", "18:00", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "10:00", "18:00", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:00", "18:00", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "10:00", "18:00", "",      "",      "Continuo"),
            ("Sábado",    "Sí", "10:00", "13:00", "",      "",      "Solo mañana"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Depilación láser",          "categoria": "Depilación"},
            {"nombre": "Manicura",                  "categoria": "Estilo personal"},
            {"nombre": "Tratamientos faciales",     "categoria": "Belleza facial"},
            {"nombre": "Limpieza facial",           "categoria": "Belleza facial"},
            {"nombre": "Medicina estética facial",  "categoria": "Medicina estética"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/beauty-by-patricia/",
    },
    {
        "nombre": "Centro de Estética Alba Guadalajara",
        "slogan": "Estética profesional con marcas de lujo",
        "descripcion": "Centro de estética ubicado en el barrio Amistad de Guadalajara. Trabaja con marcas profesionales como Natura Bissé, María Galland y Sothys. Equipa tecnología avanzada: Indiba, LPG, cavitación y presoterapia.",
        "categoria": "Estéticas",
        "especialidades": "Limpiezas faciales profundas, peeling químico, tratamientos faciales con Indiba / LPG / cavitación / presoterapia; marcas: Natura Bissé, María Galland, Sothys",
        "precio": "",
        "rating": "4.8",
        "resenas": 66,
        "direccion": "C. de Rufino Blanco, 9, 19002 Guadalajara",
        "telefono": "949 21 02 37",
        "whatsapp": "949 21 02 37",
        "email": "",
        "web": "https://www.albaesteticaguadalajara.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "19:00", "",      "",      "Continuo"),
            ("Martes",    "Sí", "10:00", "19:00", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "10:00", "19:00", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:00", "19:00", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "10:00", "19:00", "",      "",      "Continuo"),
            ("Sábado",    "Sí", "10:00", "14:00", "",      "",      "Solo mañana"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Limpieza facial profunda",  "categoria": "Belleza facial"},
            {"nombre": "Peeling químico",           "categoria": "Belleza facial"},
            {"nombre": "Tratamiento con Indiba",    "categoria": "Tratamientos avanzados"},
            {"nombre": "Tratamiento con LPG",       "categoria": "Belleza corporal"},
            {"nombre": "Cavitación",                "categoria": "Belleza corporal"},
            {"nombre": "Presoterapia",              "categoria": "Bienestar"},
            {"nombre": "Tratamientos Natura Bissé", "categoria": "Belleza facial"},
            {"nombre": "Tratamientos María Galland","categoria": "Belleza facial"},
            {"nombre": "Tratamientos Sothys",       "categoria": "Belleza facial"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/centro-de-estetica-alba-guadalajara/",
    },
    {
        "nombre": "Centro Estética Avanzada Cleó",
        "slogan": "Estética avanzada en Las Cañas",
        "descripcion": "Centro de estética avanzada ubicado en el barrio de Las Cañas, Guadalajara. Ofrece depilación láser, tratamientos faciales de última generación y técnicas como ondas de choque y viora.",
        "categoria": "Estéticas",
        "especialidades": "Depilación láser, limpieza facial, tratamientos de piel, ondas de choque, cápsula facial BDR, slimfit, viora, presoterapia, depilación con cera, perfilado de cejas con hilo",
        "precio": "",
        "rating": "4.8",
        "resenas": 49,
        "direccion": "C. Federica Montseny, 4, 19005 Guadalajara",
        "telefono": "949 03 38 94",
        "whatsapp": "949 03 38 94",
        "email": "",
        "web": "https://www.esteticacleo.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "13:30", "16:30", "20:00", ""),
            ("Martes",    "Sí", "10:00", "13:30", "16:30", "20:00", ""),
            ("Miércoles", "Sí", "10:00", "13:30", "16:30", "20:00", ""),
            ("Jueves",    "Sí", "10:00", "13:30", "16:30", "20:00", ""),
            ("Viernes",   "Sí", "10:00", "13:30", "16:30", "20:00", ""),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Depilación láser",              "categoria": "Depilación"},
            {"nombre": "Limpieza facial",               "categoria": "Belleza facial"},
            {"nombre": "Tratamientos de piel",          "categoria": "Belleza facial"},
            {"nombre": "Ondas de choque",               "categoria": "Tratamientos avanzados"},
            {"nombre": "Cápsula facial BDR",            "categoria": "Tratamientos avanzados"},
            {"nombre": "Slimfit",                       "categoria": "Belleza corporal"},
            {"nombre": "Viora",                         "categoria": "Tratamientos avanzados"},
            {"nombre": "Presoterapia",                  "categoria": "Bienestar"},
            {"nombre": "Depilación con cera",           "categoria": "Depilación"},
            {"nombre": "Perfilado de cejas con hilo",   "categoria": "Diseño de la mirada"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/centro-estetica-avanzada-cleo/",
    },
    {
        "nombre": "Clínica Carlos Gómez",
        "slogan": "Medicina estética para resultados naturales",
        "descripcion": "Clínica de medicina estética facial ubicada en el barrio Centro de Guadalajara. Especializada en tratamientos con ácido hialurónico, toxina botulínica y técnicas de rejuvenecimiento.",
        "categoria": "Medicina estética",
        "especialidades": "Depilación láser, ácido hialurónico, toxina botulínica (Botox), tratamientos con colágeno, lipolaser, Ellanse (estimulación de colágeno)",
        "precio": "",
        "rating": "4.7",
        "resenas": 92,
        "direccion": "C. Virgen del Amparo, 44, 19002 Guadalajara",
        "telefono": "949 25 73 28",
        "whatsapp": "949 25 73 28",
        "email": "",
        "web": "https://www.clinicacarlosgomez.com",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Martes",    "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Depilación láser",               "categoria": "Depilación"},
            {"nombre": "Ácido hialurónico",              "categoria": "Medicina estética"},
            {"nombre": "Toxina botulínica (Botox)",      "categoria": "Medicina estética"},
            {"nombre": "Tratamientos con colágeno",      "categoria": "Medicina estética"},
            {"nombre": "Lipolaser",                      "categoria": "Medicina estética"},
            {"nombre": "Ellanse (estimulación colágeno)","categoria": "Medicina estética"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/clinica-carlos-gomez/",
    },
    {
        "nombre": "GN Estética Avanzada",
        "slogan": "Tratamientos faciales de vanguardia",
        "descripcion": "Centro de estética avanzada ubicado en el barrio de Los Valles, Guadalajara. Especializado en tratamientos de medicina estética con tecnología de última generación.",
        "categoria": "Estéticas",
        "especialidades": "Maderoterapia con masajes drenantes, limpieza facial, mesoterapia inyectada, dermapen, Hollywood peel, plasma, hidrataciones con vitaminas, tratamientos para acné, reducción de papada",
        "precio": "",
        "rating": "4.7",
        "resenas": 24,
        "direccion": "Av. Bulevar del Alto Tajo, 9, local 4, 19005 Guadalajara",
        "telefono": "624 40 49 09",
        "whatsapp": "624 40 49 09",
        "email": "",
        "web": "https://www.gnesteticaavanzada.com",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:30", "20:00", "",      "",      "Continuo"),
            ("Martes",    "Sí", "10:30", "20:00", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "10:30", "20:00", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:30", "20:00", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "10:30", "20:00", "",      "",      "Continuo"),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Maderoterapia con masaje drenante",  "categoria": "Bienestar"},
            {"nombre": "Limpieza facial",                    "categoria": "Belleza facial"},
            {"nombre": "Mesoterapia inyectada",              "categoria": "Medicina estética"},
            {"nombre": "Dermapen",                           "categoria": "Tratamientos avanzados"},
            {"nombre": "Hollywood peel",                     "categoria": "Belleza facial"},
            {"nombre": "Plasma",                             "categoria": "Medicina estética"},
            {"nombre": "Hidrataciones con vitaminas",        "categoria": "Belleza facial"},
            {"nombre": "Tratamiento para acné",              "categoria": "Belleza facial"},
            {"nombre": "Reducción de papada",                "categoria": "Medicina estética"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/gn-estetica-avanzada/",
    },
    {
        "nombre": "Samira Salud y Belleza",
        "slogan": "Salud, belleza y bienestar en un solo centro",
        "descripcion": "Centro integral de salud y belleza ubicado en el barrio de Constitución, Guadalajara. Combina estética avanzada, depilación láser, peluquería y pilates en un mismo espacio.",
        "categoria": "Estéticas",
        "especialidades": "Tratamientos faciales (punta diamante, LPG, Indiba), depilación láser, lifting de pestañas, pilates, peluquería",
        "precio": "",
        "rating": "4.6",
        "resenas": 327,
        "direccion": "Av. del Ejército, 9A, 19004 Guadalajara",
        "telefono": "949 25 53 79",
        "whatsapp": "949 25 53 79",
        "email": "",
        "web": "https://www.samirasaludybelleza.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Martes",    "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "10:00", "20:00", "",      "",      "Continuo"),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Tratamiento facial punta de diamante",  "categoria": "Belleza facial"},
            {"nombre": "Tratamiento con LPG",                   "categoria": "Belleza facial/corporal"},
            {"nombre": "Tratamiento con Indiba",                "categoria": "Tratamientos avanzados"},
            {"nombre": "Depilación láser",                      "categoria": "Depilación"},
            {"nombre": "Lifting de pestañas",                   "categoria": "Diseño de la mirada"},
            {"nombre": "Pilates",                               "categoria": "Bienestar"},
            {"nombre": "Peluquería",                            "categoria": "Estilo personal"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/samira-salud-y-belleza/",
    },
    {
        "nombre": "Laude Servibelleza",
        "slogan": "Tu centro de estética en Entrepeñas",
        "descripcion": "Centro de estética facial ubicado en Aguas Vivas, Guadalajara. Especializado en limpiezas faciales, peelings, dermapen y tratamientos corporales.",
        "categoria": "Estéticas",
        "especialidades": "Tratamientos faciales, limpiezas faciales, peelings, dermapen, venta de serums y cremas, eliminación de quistes, pedicura",
        "precio": "~65-69 € (tratamientos según directorio)",
        "rating": "4.6",
        "resenas": 153,
        "direccion": "Blvr. de Entrepeñas, 8, 19005 Guadalajara",
        "telefono": "640 08 40 04",
        "whatsapp": "640 08 40 04",
        "email": "",
        "web": "https://www.servibelleza.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "09:30", "19:30", "",      "",      "Continuo"),
            ("Martes",    "Sí", "09:30", "19:30", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "09:30", "19:30", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "09:30", "19:30", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "09:30", "19:30", "",      "",      "Continuo"),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Tratamientos faciales",        "categoria": "Belleza facial"},
            {"nombre": "Limpiezas faciales",           "categoria": "Belleza facial"},
            {"nombre": "Peeling",                      "categoria": "Belleza facial"},
            {"nombre": "Dermapen",                     "categoria": "Tratamientos avanzados"},
            {"nombre": "Venta de serums y cremas",     "categoria": "Productos"},
            {"nombre": "Eliminación de quistes",       "categoria": "Medicina estética"},
            {"nombre": "Pedicura",                     "categoria": "Estilo personal"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/laude-servibelleza/",
    },
    {
        "nombre": "GLOW Medicina Estética",
        "slogan": "Tu piel en manos expertas",
        "descripcion": "Clínica de medicina estética ubicada en el barrio de Amistad, Guadalajara. Especializada en tratamiento de manchas faciales y depilación láser.",
        "categoria": "Medicina estética",
        "especialidades": "Tratamiento de manchas faciales, depilación láser, medicina estética facial",
        "precio": "",
        "rating": "4.3",
        "resenas": 18,
        "direccion": "C. de la Virgen de la Soledad, 3a, 19003 Guadalajara",
        "telefono": "949 21 11 72",
        "whatsapp": "949 21 11 72",
        "email": "",
        "web": "https://www.medicinaesteticaglow.es",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "?",  "",      "",      "",      "",      "Horario no publicado"),
            ("Martes",    "?",  "",      "",      "",      "",      "Horario no publicado"),
            ("Miércoles", "?",  "",      "",      "",      "",      "Horario no publicado"),
            ("Jueves",    "?",  "",      "",      "",      "",      "Horario no publicado"),
            ("Viernes",   "?",  "",      "",      "",      "",      "Horario no publicado"),
            ("Sábado",    "?",  "",      "",      "",      "",      "Verificar"),
            ("Domingo",   "?",  "",      "",      "",      "",      "Verificar"),
        ],
        "servicios": [
            {"nombre": "Tratamiento de manchas faciales",  "categoria": "Belleza facial"},
            {"nombre": "Depilación láser",                 "categoria": "Depilación"},
            {"nombre": "Medicina estética facial",         "categoria": "Medicina estética"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/glow-medicina-estetica/",
    },
    {
        "nombre": "Mizu Guadalajara",
        "slogan": "Spa y medicina estética en el corazón de Guadalajara",
        "descripcion": "Centro de spa y medicina estética ubicado en el barrio de Amistad, Guadalajara. Ofrece una amplia gama de servicios desde masajes en pareja hasta tratamientos faciales avanzados y fisioterapia estética.",
        "categoria": "Estéticas",
        "especialidades": "Medicina estética facial, masajes (relajantes, en pareja, con aromaterapia), tratamientos faciales para acné, láser estético, fisioterapia estética, cama de agua, manicura, limpieza de cutis, bonos regalo",
        "precio": "",
        "rating": "4.2",
        "resenas": 349,
        "direccion": "C. Marqués de Santillana, 2, 19002 Guadalajara",
        "telefono": "685 53 21 67",
        "whatsapp": "685 53 21 67",
        "email": "",
        "web": "https://www.mizuspa.com",
        "instagram": "",
        "facebook": "",
        "horario": [
            ("Lunes",     "Sí", "10:00", "20:30", "",      "",      "Continuo"),
            ("Martes",    "Sí", "10:00", "20:30", "",      "",      "Continuo"),
            ("Miércoles", "Sí", "10:00", "20:30", "",      "",      "Continuo"),
            ("Jueves",    "Sí", "10:00", "20:30", "",      "",      "Continuo"),
            ("Viernes",   "Sí", "10:00", "20:30", "",      "",      "Continuo"),
            ("Sábado",    "No", "",      "",      "",      "",      "Cerrado"),
            ("Domingo",   "No", "",      "",      "",      "",      "Cerrado"),
        ],
        "servicios": [
            {"nombre": "Medicina estética facial",          "categoria": "Medicina estética"},
            {"nombre": "Masajes relajantes",                "categoria": "Bienestar"},
            {"nombre": "Masajes en pareja",                 "categoria": "Bienestar"},
            {"nombre": "Masajes con aromaterapia",          "categoria": "Bienestar"},
            {"nombre": "Tratamientos faciales para acné",   "categoria": "Belleza facial"},
            {"nombre": "Láser estético",                    "categoria": "Tratamientos avanzados"},
            {"nombre": "Fisioterapia estética",             "categoria": "Bienestar"},
            {"nombre": "Cama de agua",                      "categoria": "Bienestar"},
            {"nombre": "Manicura",                          "categoria": "Estilo personal"},
            {"nombre": "Limpieza de cutis",                 "categoria": "Belleza facial"},
            {"nombre": "Bonos regalo",                      "categoria": "Productos"},
        ],
        "fuente_url": "https://mejoresclinicas.net/clinica/mizu-guadalajara/",
    },
]

# ─── Generar todos los Excel ──────────────────────────────────────────────────
print(f"\n{'='*60}")
print(f"  Generando fichas Fluix — {len(NEGOCIOS)} negocios")
print(f"  Destino: {OUTPUT_DIR}")
print(f"{'='*60}\n")

for negocio in NEGOCIOS:
    crear_excel(negocio)

print(f"\n{'='*60}")
print(f"  OK {len(NEGOCIOS)} fichas generadas en {OUTPUT_DIR}")
print(f"{'='*60}\n")
